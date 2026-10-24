import asyncio
import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
import logging
from pathlib import PurePath
import tempfile
import os

from google import genai
from google.genai import types
from pypdf import PdfReader

from services.chunking import normalize_text
from settings import (
    LLM_MODEL,
    MAX_PDF_PAGES,
    MAX_VIDEO_MINUTES,
    MAX_INDEX_FILE_BYTES,
)

logger = logging.getLogger("spaces.extraction")


class UnsupportedDocumentType(ValueError):
    pass


class DocumentLimitExceeded(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedUnit:
    location: str  # e.g. "page 3", "slide 12", "image", "00:12:30"
    content: str


@dataclass(frozen=True)
class ExtractedDocument:
    status: str  # "READY", "FAILED", "SKIPPED"
    units: list[ExtractedUnit]
    error: str | None = None


async def extract_document(
    file_name: str,
    content: bytes,
    client: genai.Client | None = None,
) -> ExtractedDocument:
    extension = PurePath(file_name).suffix.lower()

    # 1. File size check
    if len(content) > MAX_INDEX_FILE_BYTES:
        raise DocumentLimitExceeded(
            f"File size ({len(content) // (1024*1024)}MB) exceeds maximum limit of {MAX_INDEX_FILE_BYTES // (1024*1024)}MB."
        )

    # -------------------------------------------------------------------------
    # PDF Extraction (with OCR fallback for scanned pages)
    # -------------------------------------------------------------------------
    if extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(content), strict=False)
            num_pages = len(reader.pages)
            if num_pages > MAX_PDF_PAGES:
                raise DocumentLimitExceeded(
                    f"PDF page count ({num_pages}) exceeds maximum allowed ({MAX_PDF_PAGES} pages)."
                )

            units: list[ExtractedUnit] = []
            total_text_len = 0

            for page_idx, page in enumerate(reader.pages, start=1):
                raw_text = page.extract_text() or ""
                norm_text = normalize_text(raw_text)
                if norm_text:
                    total_text_len += len(norm_text)
                    units.append(ExtractedUnit(location=f"page {page_idx}", content=norm_text))

            # Scanned PDF detection: very little or no extractable text
            if total_text_len < 60 and client:
                logger.info(f"PDF {file_name} appears scanned ({total_text_len} chars). Invoking Gemini OCR...")
                try:
                    ocr_response = await client.aio.models.generate_content(
                        model=LLM_MODEL,
                        contents=[
                            types.Part.from_bytes(data=content, mime_type="application/pdf"),
                            "Extract all visible text from this scanned document verbatim page by page. "
                            "Prefix each page with [Page X] and preserve all tables, numbers, and headings.",
                        ],
                    )
                    ocr_text = ocr_response.text or ""
                    if ocr_text.strip():
                        # Split by [Page X] headers if possible
                        import re
                        page_splits = re.split(r"\[Page\s*(\d+)\]", ocr_text, flags=re.IGNORECASE)
                        if len(page_splits) > 1:
                            ocr_units = []
                            for i in range(1, len(page_splits), 2):
                                p_num = page_splits[i]
                                p_txt = normalize_text(page_splits[i + 1])
                                if p_txt:
                                    ocr_units.append(ExtractedUnit(location=f"page {p_num}", content=p_txt))
                            if ocr_units:
                                return ExtractedDocument(status="READY", units=ocr_units)
                        return ExtractedDocument(
                            status="READY",
                            units=[ExtractedUnit(location="scanned document", content=normalize_text(ocr_text))],
                        )
                except Exception as ocr_err:
                    logger.warning(f"OCR fallback failed: {ocr_err}")

            if not units:
                return ExtractedDocument(
                    status="FAILED",
                    units=[],
                    error="No readable text found in PDF document.",
                )
            return ExtractedDocument(status="READY", units=units)

        except DocumentLimitExceeded:
            raise
        except Exception as error:
            raise ValueError(f"The PDF could not be read: {error}") from error

    # -------------------------------------------------------------------------
    # PPTX Presentation Extraction (Slide Text + Speaker Notes)
    # -------------------------------------------------------------------------
    if extension in {".pptx", ".ppt"}:
        try:
            from pptx import Presentation

            prs = Presentation(BytesIO(content))
            units = []

            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_parts: list[str] = []

                # Extract shapes text & tables
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            line = paragraph.text.strip()
                            if line:
                                slide_parts.append(line)
                    elif shape.has_table and shape.table:
                        for row in shape.table.rows:
                            row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                            if row_txt:
                                slide_parts.append(f"| {row_txt} |")

                # Extract speaker notes
                if slide.has_notes_slide and slide.notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame
                    if notes_frame and notes_frame.text:
                        notes_clean = notes_frame.text.strip()
                        if notes_clean:
                            slide_parts.append(f"[Speaker Notes]: {notes_clean}")

                slide_text = normalize_text("\n".join(slide_parts))
                if slide_text:
                    units.append(ExtractedUnit(location=f"slide {slide_idx}", content=slide_text))

            if not units:
                return ExtractedDocument(
                    status="FAILED",
                    units=[],
                    error="Presentation has no readable text or notes on slides.",
                )
            return ExtractedDocument(status="READY", units=units)
        except Exception as error:
            raise ValueError(f"The presentation could not be read: {error}") from error

    # -------------------------------------------------------------------------
    # DOCX Document Extraction (Headings + Paragraphs + Tables)
    # -------------------------------------------------------------------------
    if extension in {".docx", ".doc"}:
        try:
            import docx

            doc = docx.Document(BytesIO(content))
            units = []
            current_section = "section 1"
            section_idx = 1
            section_paragraphs: list[str] = []

            for element in doc.paragraphs:
                txt = element.text.strip()
                if not txt:
                    continue
                # Heading detection
                if element.style and element.style.name and element.style.name.startswith("Heading"):
                    if section_paragraphs:
                        units.append(
                            ExtractedUnit(
                                location=current_section,
                                content=normalize_text("\n\n".join(section_paragraphs)),
                            )
                        )
                        section_paragraphs = []
                        section_idx += 1
                    current_section = txt if len(txt) <= 50 else f"section {section_idx}"
                section_paragraphs.append(txt)

            # Extract tables as markdown tables
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        table_rows.append(f"| {row_txt} |")
                if table_rows:
                    section_paragraphs.append("\n".join(table_rows))

            if section_paragraphs:
                units.append(
                    ExtractedUnit(
                        location=current_section,
                        content=normalize_text("\n\n".join(section_paragraphs)),
                    )
                )

            if not units:
                return ExtractedDocument(status="FAILED", units=[], error="DOCX document contains no text.")
            return ExtractedDocument(status="READY", units=units)
        except Exception as error:
            raise ValueError(f"The DOCX document could not be read: {error}") from error

    # -------------------------------------------------------------------------
    # Plain Text / Markdown / CSV
    # -------------------------------------------------------------------------
    if extension in {".txt", ".md", ".markdown"}:
        text = normalize_text(content.decode("utf-8-sig", errors="replace"))
        if not text:
            return ExtractedDocument(status="FAILED", units=[], error="Text document is empty.")
        return ExtractedDocument(status="READY", units=[ExtractedUnit(location="document", content=text)])

    if extension == ".csv":
        try:
            decoded = content.decode("utf-8-sig", errors="replace")
            reader = csv.reader(StringIO(decoded))
            rows = list(reader)
            if not rows:
                return ExtractedDocument(status="FAILED", units=[], error="CSV is empty.")

            # Chunk CSV in batches of 40 rows preserving header
            header = " | ".join(rows[0])
            units = []
            chunk_rows = []
            batch_num = 1

            for row in rows[1:]:
                chunk_rows.append(" | ".join(row))
                if len(chunk_rows) >= 40:
                    table_md = f"| {header} |\n" + "\n".join(f"| {r} |" for r in chunk_rows)
                    units.append(ExtractedUnit(location=f"rows {batch_num*40-39}-{batch_num*40}", content=table_md))
                    chunk_rows = []
                    batch_num += 1

            if chunk_rows:
                table_md = f"| {header} |\n" + "\n".join(f"| {r} |" for r in chunk_rows)
                units.append(ExtractedUnit(location=f"rows batch {batch_num}", content=table_md))

            return ExtractedDocument(status="READY", units=units)
        except Exception as error:
            raise ValueError(f"CSV could not be read: {error}") from error

    # -------------------------------------------------------------------------
    # Image Transcription & Description (Gemini Vision)
    # -------------------------------------------------------------------------
    if extension in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        if not client:
            raise ValueError("Gemini client is required to process image files.")

        mime_type = "image/jpeg"
        if extension == ".png":
            mime_type = "image/png"
        elif extension == ".webp":
            mime_type = "image/webp"
        elif extension == ".gif":
            mime_type = "image/gif"

        try:
            prompt = (
                "Transcribe all visible text in this image verbatim. "
                "Then provide a concise description of the image content including: "
                "objects, people count (if any), setting/place, and main activity or diagram concept."
            )
            response = await client.aio.models.generate_content(
                model=LLM_MODEL,
                contents=[
                    types.Part.from_bytes(data=content, mime_type=mime_type),
                    prompt,
                ],
            )
            extracted = normalize_text(response.text or "")
            if not extracted:
                return ExtractedDocument(
                    status="FAILED",
                    units=[],
                    error="Could not extract text or description from image.",
                )
            return ExtractedDocument(
                status="READY",
                units=[ExtractedUnit(location="image", content=extracted)],
            )
        except Exception as error:
            raise ValueError(f"Gemini image analysis failed: {error}") from error

    # -------------------------------------------------------------------------
    # Video & Audio (Gemini Files API with Timestamps)
    # -------------------------------------------------------------------------
    if extension in {".mp4", ".webm", ".mov", ".mp3", ".wav", ".m4a"}:
        if not client:
            raise ValueError("Gemini client is required to process video/audio files.")

        mime_map = {
            ".mp4": "video/mp4",
            ".webm": "video/webm",
            ".mov": "video/quicktime",
            ".mp3": "audio/mp3",
            ".wav": "audio/wav",
            ".m4a": "audio/m4a",
        }
        mime_type = mime_map.get(extension, "video/mp4")

        # Write to temporary file for Files API upload
        temp_file = tempfile.NamedTemporaryFile(suffix=extension, delete=False)
        try:
            temp_file.write(content)
            temp_file.flush()
            temp_file.close()

            logger.info(f"Uploading media {file_name} to Gemini Files API...")
            uploaded = await asyncio.to_thread(
                client.files.upload,
                file=temp_file.name,
                mime_type=mime_type,
            )

            # Wait until state is ACTIVE
            max_wait = 120  # seconds
            elapsed = 0
            while uploaded.state.name == "PROCESSING" and elapsed < max_wait:
                await asyncio.sleep(4)
                elapsed += 4
                uploaded = await asyncio.to_thread(client.files.get, name=uploaded.name)

            if uploaded.state.name != "ACTIVE":
                raise ValueError(f"Media processing failed in Gemini Files API with state: {uploaded.state.name}")

            prompt = (
                "Generate a detailed transcript of this media with timestamps (format: [HH:MM:SS]) "
                "for each spoken segment, along with a short on-screen visual description of key scenes. "
                "Break into chronological segments."
            )
            response = await client.aio.models.generate_content(
                model=LLM_MODEL,
                contents=[uploaded, prompt],
            )

            # Clean up uploaded file from Gemini storage
            try:
                await asyncio.to_thread(client.files.delete, name=uploaded.name)
            except Exception:
                pass

            raw_transcript = response.text or ""
            if not raw_transcript.strip():
                return ExtractedDocument(
                    status="FAILED",
                    units=[],
                    error="No audible speech or transcript generated for this media.",
                )

            # Parse timestamps like [00:01:23] into ExtractedUnits
            import re
            timestamp_blocks = re.split(r"\[(\d{1,2}:\d{2}(?::\d{2})?)\]", raw_transcript)
            units = []
            if len(timestamp_blocks) > 1:
                for i in range(1, len(timestamp_blocks), 2):
                    ts = timestamp_blocks[i]
                    segment_text = normalize_text(timestamp_blocks[i + 1])
                    if segment_text:
                        units.append(ExtractedUnit(location=f"timestamp {ts}", content=segment_text))

            if not units:
                units = [ExtractedUnit(location="timestamp 00:00:00", content=normalize_text(raw_transcript))]

            return ExtractedDocument(status="READY", units=units)

        finally:
            if os.path.exists(temp_file.name):
                try:
                    os.unlink(temp_file.name)
                except Exception:
                    pass

    raise UnsupportedDocumentType(
        f"Unsupported file format '{extension}'. Supported formats: PDF, PPTX, DOCX, TXT, MD, CSV, Images, Audio/Video."
    )